"""Local, deterministic document templates. No LLM invents fields or legal citations."""
import io
import os
from pathlib import Path
from xml.sax.saxutils import escape

from app.cases import Case
from app.schemas import Draft

TITLES = {
    "police": ("Заявление о предполагаемом мошенничестве", "Ықтимал алаяқтық туралы арыз"),
    "bank_dispute": ("Обращение о спорной операции", "Даулы операция туралы өтініш"),
    "bank_loan": ("Обращение о спорном займе", "Даулы қарыз туралы өтініш"),
    "regulator": ("Обращение о действиях финансовой организации", "Қаржы ұйымының әрекеттері туралы өтініш"),
    "block_reason": ("Запрос об основаниях ограничения счёта", "Шотты шектеу негіздері туралы сұрау"),
    "restructure": ("Просьба рассмотреть временные меры по займу", "Қарыз бойынша уақытша шараларды қарау туралы өтініш"),
    "evidence": ("Памятка по сохранению доказательств", "Дәлелдерді сақтау жөніндегі жадынама"),
}
REQUESTS = {
    "police": ("Прошу принять и зарегистрировать моё обращение, проверить изложенные обстоятельства и сообщить регистрационные данные и результат рассмотрения.", "Арызымды қабылдап, тіркеуді, көрсетілген мән-жайларды тексеруді және тіркеу деректері мен қарау нәтижесін хабарлауды сұраймын."),
    "bank_dispute": ("Прошу проверить указанную операцию, рассмотреть доступные меры защиты и возврата средств, предоставить сведения об операции и письменный ответ.", "Көрсетілген операцияны тексеруді, қорғау және қаражатты қайтарудың қолжетімді шараларын қарауды, операция туралы мәлімет пен жазбаша жауап беруді сұраймын."),
    "bank_loan": ("Оспариваю обстоятельства оформления указанного займа. Прошу провести проверку, предоставить копию договора и сведения о способе его оформления, рассмотреть временные защитные меры и дать письменный ответ.", "Көрсетілген қарыздың ресімделу мән-жайларына дау білдіремін. Тексеру жүргізуді, шарт көшірмесін және ресімдеу тәсілі туралы мәлімет беруді, уақытша қорғау шараларын қарауды және жазбаша жауап беруді сұраймын."),
    "regulator": ("Прошу рассмотреть действия финансовой организации в пределах компетенции Агентства, проверить изложенные обстоятельства и предоставить мотивированный ответ.", "Қаржы ұйымының әрекеттерін Агенттік құзыреті шегінде қарауды, көрсетілген мән-жайларды тексеруді және дәлелді жауап беруді сұраймын."),
    "block_reason": ("Прошу письменно сообщить основания, дату и объём ограничения, реквизиты решения и сведения об органе, если ограничение введено по его решению, а также порядок дальнейшего обращения и перечень необходимых документов.", "Шектеудің негіздерін, күнін және көлемін, шешім деректемелерін және оны қабылдаған орган туралы мәліметті, сондай-ақ әрі қарай жүгіну тәртібі мен қажетті құжаттар тізімін жазбаша хабарлауды сұраймын."),
    "restructure": ("Прошу рассмотреть возможность временных мер, отсрочки или изменения графика на период выяснения обстоятельств и письменно сообщить условия и последствия предлагаемых вариантов. Это обращение не содержит согласия с оспариваемыми обстоятельствами возникновения долга.", "Мән-жайлар анықталған кезеңде уақытша шаралар, төлемді кейінге қалдыру немесе кестені өзгерту мүмкіндігін қарауды және ұсынылатын нұсқалардың шарттары мен салдарын жазбаша хабарлауды сұраймын. Бұл өтініш даулы берешектің туындау мән-жайларымен келісуді білдірмейді."),
}


def build_draft(kind: str, case: Case, language: str) -> Draft:
    if kind not in TITLES or not case.confirmed:
        raise ValueError("Confirm facts before generating documents")
    kk = language == "kk"
    if kind == "evidence":
        text = ("Сохраните исходную переписку, выписки, чеки и договоры. Зафиксируйте время событий и адреса страниц. "
                "Не удаляйте и не редактируйте оригиналы. Передавайте копии только через проверенные каналы.\n\n" if not kk else
                "Бастапқы хаттарды, үзінді көшірмелерді, түбіртектер мен шарттарды сақтаңыз. Оқиға уақытын және беттердің мекенжайларын белгілеңіз. "
                "Түпнұсқаларды жоймаңыз және өзгертпеңіз. Көшірмелерді тек тексерілген арналар арқылы беріңіз.\n\n")
        return Draft(title=TITLES[kind][kk], body=text + case.summary(language), missing_fields=[], attachments=[])
    recipient = ("[ТОЛТЫРЫҢЫЗ: адресат]" if kk else "[ЗАПОЛНИТЬ: адресат]")
    if kind == "regulator":
        recipient = ("Қазақстан Республикасының Қаржы нарығын реттеу және дамыту агенттігіне" if kk else
                     "В Агентство Республики Казахстан по регулированию и развитию финансового рынка")
    sender = "[ТОЛТЫРЫҢЫЗ: өтініш беруші және байланыс]" if kk else "[ЗАПОЛНИТЬ: заявитель и контакт для ответа]"
    appendix = "[ТОЛТЫРЫҢЫЗ: нақты тіркелетін құжаттар]" if kk else "[ЗАПОЛНИТЬ: фактически прилагаемые документы]"
    ending = "Күні: [ТОЛТЫРЫҢЫЗ]\nҚолы: [ТОЛТЫРЫҢЫЗ]" if kk else "Дата: [ЗАПОЛНИТЬ]\nПодпись: [ЗАПОЛНИТЬ]"
    body = f"{recipient}\n{sender}\n\n{case.summary(language)}\n\n{REQUESTS[kind][kk]}\n\n{appendix}\n\n{ending}"
    return Draft(title=TITLES[kind][kk], body=body, missing_fields=[sender, appendix, ending], attachments=[])


def font_file(custom: str = "") -> Path:
    candidates = [custom, os.environ.get("PDF_FONT_PATH", ""),
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise RuntimeError("A Cyrillic/Kazakh TrueType font is required")


def export(draft: Draft, fmt: str, font_path: str = "") -> bytes:
    if fmt == "txt":
        return (draft.title + "\n\n" + draft.body).encode("utf-8")
    buffer = io.BytesIO()
    if fmt == "docx":
        from docx import Document
        from docx.shared import Mm, Pt, RGBColor
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(22)
        for name in ["Normal", "Title"]:
            doc.styles[name].font.name = "Arial"
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles["Normal"].font.size = Pt(11)
        doc.styles["Title"].font.size = Pt(16)
        doc.add_paragraph(draft.title, "Title")
        for line in draft.body.split("\n"):
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)
        doc.save(buffer)
    elif fmt == "pdf":
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        name = "ZanUnicode"
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_file(font_path))))
        style = ParagraphStyle("body", fontName=name, fontSize=10.5, leading=15, alignment=TA_LEFT,
                               spaceAfter=6, splitLongWords=True)
        title = ParagraphStyle("title", parent=style, fontSize=15, leading=20, spaceAfter=18)
        story = [Paragraph(escape(draft.title), title)]
        for line in draft.body.split("\n"):
            story.append(Paragraph(escape(line), style) if line else Spacer(1, 7))
        SimpleDocTemplate(buffer, pagesize=(210*mm, 297*mm), leftMargin=22*mm, rightMargin=22*mm,
                          topMargin=20*mm, bottomMargin=20*mm, title=draft.title, author="Zan Protect").build(story)
    else:
        raise ValueError("Unsupported output format")
    return buffer.getvalue()
