import io
import time
import zipfile

import pytest

from app.files import extract, extract_local
from app.llm import trusted_url, validate_answer
from app.privacy import Sessions, redact, session_key
from app.schemas import Answer, Claim, Intake, Source


def sample_answer(risk="low", claims=None):
    return Answer(language="ru", risk=risk, summary="test", evidence=[], explanation="test",
                  immediate_actions=[], next_actions=[], questions=[], legal_claims=claims or [],
                  uncertainty="", lawyer_needed=False)


def intake(risk):
    return Intake(language="ru", situation="check", summary="test", indicators=[],
                  risk_floor=risk, search_topic="none", questions=[])


@pytest.mark.parametrize("url", ["http://gov.kz/x", "https://gov.kz.evil.com", "https://evilgov.kz", "https://gov.kz@evil.com"])
def test_untrusted_domain(url):
    assert not trusted_url(url)


def test_official_domain():
    assert trusted_url("https://www.gov.kz/memleket/entities/ardfm")


def test_missing_citation_rejected():
    with pytest.raises(ValueError):
        validate_answer(sample_answer(claims=[Claim(text="claim", source_ids=["invented"])]), intake("low"), [])


def test_empty_citation_rejected():
    with pytest.raises(ValueError):
        validate_answer(sample_answer(claims=[Claim(text="claim", source_ids=[])]), intake("low"), [])


def test_valid_citation():
    source = Source(id="s1", title="x", url="https://gov.kz", checked_on="2026-10-11", excerpt="x")
    answer = sample_answer(claims=[Claim(text="claim", source_ids=["s1"])])
    assert validate_answer(answer, intake("low"), [source]).legal_claims


@pytest.mark.parametrize("floor,expected", [("high", "high"), ("medium", "medium"), ("unknown", "unknown")])
def test_risk_not_downgraded(floor, expected):
    assert validate_answer(sample_answer(), intake(floor), []).risk == expected


def test_pii_masking():
    text = redact("ИИН 990101123456 email test@example.com +7 701 123 45 67 https://evil.test/code?secret=123")
    for private in ["990101123456", "test@example.com", "701", "evil.test", "secret=123"]:
        assert private not in text


def test_session_expiry_and_deletion():
    sessions = Sessions(60)
    s = sessions.get("a")
    s.history.append("private")
    s.expires = time.monotonic() - 1
    sessions.purge()
    assert not sessions.items
    assert sessions.get("a").history == []
    sessions.delete("a")
    assert not sessions.items


def test_hmac_scope():
    assert session_key(123, "a") != session_key(123, "b")
    assert session_key(123, "a") != session_key(124, "a")


@pytest.mark.parametrize("name,data", [("virus.exe", b"MZ"), ("fake.pdf", b"not a pdf"), ("empty.txt", b"")])
def test_bad_files(name, data):
    with pytest.raises(ValueError):
        extract_local(name, data)


def test_archive_expansion_limit():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (51 * 1024 * 1024))
    with pytest.raises(ValueError, match="Expanded archive"):
        extract_local("bomb.docx", output.getvalue())


async def test_subprocess_text():
    result = await extract("example.txt", "Алаяқтық / мошенничество".encode())
    assert result["text"] == "Алаяқтық / мошенничество"
    assert result["image"] is None


def test_docx_tables():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Казахстан")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "300 000 ₸"
    output = io.BytesIO()
    doc.save(output)
    text = extract_local("case.docx", output.getvalue())["text"]
    assert "Казахстан" in text and "300 000" in text


def test_pptx_text():
    from pptx import Presentation
    pres = Presentation()
    pres.slides.add_slide(pres.slide_layouts[1]).shapes.title.text = "Картаңызды жалға"
    output = io.BytesIO()
    pres.save(output)
    assert "Картаңызды жалға" in extract_local("case.pptx", output.getvalue())["text"]


def test_xlsx_text():
    from openpyxl import Workbook
    wb = Workbook()
    wb.active["A1"] = "Аударым"
    output = io.BytesIO()
    wb.save(output)
    assert "Аударым" in extract_local("case.xlsx", output.getvalue())["text"]


def test_image_reencoding():
    from PIL import Image
    output = io.BytesIO()
    Image.new("RGB", (32, 32), "white").save(output, "PNG")
    assert extract_local("photo.png", output.getvalue())["image"].startswith("data:image/jpeg;base64,")


def test_text_pdf():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 250 Td (Fraud report 300000 KZT) Tj ET")
    page[NameObject("/Contents")] = stream
    output = io.BytesIO()
    writer.write(output)
    assert "300000" in extract_local("report.pdf", output.getvalue())["text"]


def test_scan_pdf_not_marked_safe():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="No text"):
        extract_local("scan.pdf", output.getvalue())
