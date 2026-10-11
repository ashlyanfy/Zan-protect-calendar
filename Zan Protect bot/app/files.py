"""Bounded extraction in a killable subprocess. No shell, macros, remote links or file persistence."""
import asyncio
import base64
import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

MAX_TEXT = 24000
MAX_BYTES = 10 * 1024 * 1024
OFFICE = {".docx": "word/", ".pptx": "ppt/", ".xlsx": "xl/"}
SUPPORTED = {".txt", ".csv", ".json", ".md", ".pdf", ".png", ".jpg", ".jpeg", ".webp", *OFFICE}


def extract_local(name: str, data: bytes) -> dict:
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED or not data or len(data) > MAX_BYTES:
        raise ValueError("Unsupported or oversized file")
    if ext in OFFICE:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(x.file_size for x in entries) > 50 * 1024 * 1024:
                raise ValueError("Expanded archive limit")
            if not any(x.filename.startswith(OFFICE[ext]) for x in entries):
                raise ValueError("Office format mismatch")
            if any("vbaproject" in x.filename.lower() or x.flag_bits & 1 for x in entries):
                raise ValueError("Macros/encryption are not supported")
    stream = io.BytesIO(data)
    text = ""
    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 16_000_000
        with Image.open(stream) as img:
            if img.width * img.height > 16_000_000:
                raise ValueError("Image dimensions exceeded")
            img.load()
            img.thumbnail((2400, 2400))
            output = io.BytesIO()
            img.convert("RGB").save(output, "JPEG", quality=85)
        return {"text": "", "image": "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode()}
    if ext == ".pdf":
        from pypdf import PdfReader
        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF signature mismatch")
        reader = PdfReader(stream)
        if reader.is_encrypted or len(reader.pages) > 40:
            raise ValueError("Encrypted PDF or more than 40 pages")
        for index, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if len(page_text.strip()) < 15:
                if len(reader.pages) > 10:
                    raise ValueError("OCR PDF limit: 10 pages")
                import pypdfium2 as pdfium
                import pytesseract
                pdf = pdfium.PdfDocument(data)
                try:
                    rendered = pdf[index].render(scale=1.5).to_pil()
                    try:
                        page_text = pytesseract.image_to_string(rendered, lang="rus+kaz+eng", timeout=12)
                    except (RuntimeError, OSError) as exc:
                        raise ValueError("No text: OCR unavailable or timed out") from exc
                finally:
                    pdf.close()
            text += f"[Page {index + 1}]\n" + page_text + "\n"
            if len(text) > MAX_TEXT:
                raise ValueError("Text limit: send a smaller excerpt")
    elif ext == ".docx":
        from docx import Document
        doc = Document(stream)
        text = "\n".join([p.text for p in doc.paragraphs] +
                         [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
    elif ext == ".pptx":
        from pptx import Presentation
        pres = Presentation(stream)
        if len(pres.slides) > 60:
            raise ValueError("Slide limit")
        parts = []
        for n, slide in enumerate(pres.slides, 1):
            parts.append(f"[Slide {n}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(" | ".join(c.text for c in r.cells) for r in shape.table.rows)
        text = "\n".join(parts)
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        try:
            parts = []
            cells = 0
            for sheet in wb:
                parts.append(f"[Sheet {sheet.title}]")
                for row in sheet.iter_rows():
                    cells += len(row)
                    if cells > 20000:
                        raise ValueError("Cell limit")
                    parts.append(" | ".join(str(c.value) for c in row if c.value is not None))
            text = "\n".join(parts)
        finally:
            wb.close()
    else:
        text = data.decode("utf-8-sig")
    if not text.strip():
        raise ValueError("No text: scanned PDF or image-only document; send separate images")
    if len(text) > MAX_TEXT:
        raise ValueError("Text limit: send a smaller excerpt")
    return {"text": text, "image": None}


def _run(name: str, data: bytes) -> dict:
    request = json.dumps({"name": name, "data": base64.b64encode(data).decode()}).encode()
    # Do not let document parsers inherit LLM, Telegram, database or encryption credentials.
    env = {key: value for key, value in os.environ.items()
           if key.upper() in {"PATH", "SYSTEMROOT", "WINDIR", "TEMP", "TMP", "LANG", "TESSDATA_PREFIX"}}
    result = subprocess.run([sys.executable, "-m", "app.files"], input=request, env=env,
                            capture_output=True, timeout=90, check=False)
    if result.returncode:
        raise ValueError("File could not be safely read")
    return json.loads(result.stdout)


async def extract(name: str, data: bytes) -> dict:
    if len(data) > MAX_BYTES:
        raise ValueError("File size limit")
    return await asyncio.to_thread(_run, name, data)


if __name__ == "__main__":
    try:
        if sys.platform != "win32":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_CPU, (60, 60))
        request = json.loads(sys.stdin.buffer.read(MAX_BYTES * 2))
        result = extract_local(request["name"], base64.b64decode(request["data"], validate=True))
        sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))
    except Exception:  # noqa: BLE001 -- parser subprocess returns only a sanitized failure code
        sys.exit(1)
